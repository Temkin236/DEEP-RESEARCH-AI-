#!/usr/bin/env python3
"""
Generate the Executive Showcase PowerPoint presentation for Deep Research AI.
Theme: Warm ivory, light cream, warm brown, and deep charcoal.
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Color Palette
    COLOR_BG = RGBColor(250, 247, 242)       # Warm Ivory #FAF7F2
    COLOR_CARD = RGBColor(255, 254, 252)     # White card #FFFEFC
    COLOR_PANEL = RGBColor(243, 234, 223)    # Light Cream #F3EADF
    COLOR_PRIMARY = RGBColor(184, 149, 118)  # Warm Brown #B89576
    COLOR_ACCENT = RGBColor(155, 120, 91)    # Medium Brown #9B785B
    COLOR_TEXT = RGBColor(51, 43, 37)        # Deep Charcoal #332B25
    COLOR_MUTED = RGBColor(117, 104, 93)     # Muted Brown #75685D
    COLOR_GREEN = RGBColor(82, 120, 92)      # Forest Green #52785C
    COLOR_BORDER = RGBColor(229, 216, 200)   # Border Beige #E5D8C8

    slides_content = [
        {
            "num": "01",
            "title": "Deep Research AI",
            "subtitle": "Autonomous Multi-Agent Research Platform with Verified Citations",
            "bullets": [
                "Full-Stack AI Architecture: LangGraph, FastAPI, ChromaDB, and React",
                "Decoupled 5-Agent Pipeline: Planning, Researching, Analyzing, Verifying, Reporting",
                "Deterministic Grounding: Shared Evidence Registry eliminates hallucinations",
                "Executive Warm-Ivory UI: Designed for high-clarity technical & market research"
            ]
        },
        {
            "num": "02",
            "title": "The Problem Statement",
            "subtitle": "Why Standard LLM Chat Assistants Fail at In-Depth Research",
            "bullets": [
                "Hallucinated Citations: LLMs fabricate academic papers, non-existent authors, and dead URLs",
                "Unbounded Monolithic Reasoning: Single prompts lack the ability to self-correct or cross-verify",
                "Superficial Web Summaries: General search tools scrape high-level snippets without deep synthesis",
                "Zero Auditability: Readers cannot distinguish grounded facts from model parameter assumptions"
            ]
        },
        {
            "num": "03",
            "title": "The Proposed Solution",
            "subtitle": "Decoupled Multi-Agent Architecture with Strict Citation Boundaries",
            "bullets": [
                "Specialized Agent Roles: Separate planning from research, analysis, and adversarial verification",
                "Shared Evidence Registry: Every retrieved passage is registered, normalized, and indexed",
                "Deterministic Citation Validator: Parses generated reports and strips any ungrounded citation",
                "Real-Time SSE Workspace: Complete transparency into agent thoughts, search queries, and findings"
            ]
        },
        {
            "num": "04",
            "title": "Technology Stack",
            "subtitle": "Production-Grade Engineering across Backend and Frontend",
            "bullets": [
                "Backend Framework: FastAPI (Python 3.12) with asynchronous worker execution",
                "Agent Orchestration: LangGraph StateGraph with conditional verification routing",
                "Model & Search Adapters: Google Gemini API (default), OpenAI adapter, DuckDuckGo & Tavily",
                "Persistence & RAG: PostgreSQL (relational), Alembic (migrations), and ChromaDB (vectors)",
                "Frontend Stack: React 18, Vite, TypeScript, Tailwind CSS, TanStack Query, and Lucide"
            ]
        },
        {
            "num": "05",
            "title": "System Architecture Diagram",
            "subtitle": "End-to-End Component Interaction and Egress Security",
            "bullets": [
                "Client Layer: React SPA subscribing to Server-Sent Events (SSE) with reconnect resilience",
                "API Gateway: FastAPI routing /api/v1/research and /api/v1/documents",
                "Worker Boundary: Async LangGraph stateful execution in worker threads",
                "Security Egress Layer: Strict SSRF network filter blocking private subnets (RFC-1918) and metadata IPs",
                "Data Layer: SQLAlchemy models backed by PostgreSQL / SQLite + ChromaDB vector persistence"
            ]
        },
        {
            "num": "06",
            "title": "The 5-Agent Multi-Agent Workflow",
            "subtitle": "Structured Collaboration Engineered in LangGraph",
            "bullets": [
                "1. Research Planner: Dissects inquiries into 3-4 subquestions and targeted queries",
                "2. Researcher Agent: Gathers live web snippets and queries ChromaDB for user documents",
                "3. Analyst Agent: Evaluates evidence, isolates consensus, and tags conflicting trade-offs",
                "4. Verification Agent: Fact-checks assertions against evidence registry with conditional retry loop",
                "5. Report Writer Agent: Compiles publication-quality markdown report with numbered citations"
            ]
        },
        {
            "num": "07",
            "title": "Evidence Registry & Citation Engine",
            "subtitle": "Eliminating Hallucinations Through Deterministic Grounding",
            "bullets": [
                "Source Registration: URLs are normalized, deduplicated, and given persistent indices [1], [2]",
                "Claim-to-Evidence Association: Every finding explicitly links back to excerpt source IDs",
                "Citation Validator: Scans report text, ensures every bracketed index exists in the registry",
                "Automatic Pruning: Hallucinated or non-existent citations are stripped or flagged before render"
            ]
        },
        {
            "num": "08",
            "title": "Document RAG Integration",
            "subtitle": "Fusing Internal Private Knowledge with Live Web Search",
            "bullets": [
                "Multi-Format Ingestion: Supports PDF, DOCX, Markdown (.md), and plain text (.txt)",
                "Recursive Chunking: 1,000-character windows with 200-character overlap for context coherence",
                "Vector Similarity Search: Persistent ChromaDB collection queries top-k relevant chunks",
                "Unified Evidence Pipeline: Document passages receive internal URIs and compete equally in reports"
            ]
        },
        {
            "num": "09",
            "title": "User Interface & Experience",
            "subtitle": "Warm Ivory, Cream, and Light-Brown Executive Design",
            "bullets": [
                "Executive Aesthetics: Custom HSL palette featuring warm ivory (#FAF7F2) and light brown (#B89576)",
                "Live Research Workspace: Real-time progress bar, agent timeline, and live source cards",
                "Interactive Report View: Safe Markdown rendering, clickable citation chips, and comparison tables",
                "Comprehensive Exporting: Raw Markdown copy, .md file download, and print-friendly formatting"
            ]
        },
        {
            "num": "10",
            "title": "Testing & Security Hardening",
            "subtitle": "Verified Quality Across All Tiers",
            "bullets": [
                "Backend Testing: 18 passing pytest tests covering security, agents, RAG, and API endpoints",
                "Frontend Testing: Vitest suite verifying Navbar, ProgressIndicator, and SourceCards",
                "SSRF Protection: Mandatory IP checks block loopbacks, AWS/GCP metadata IPs, and local subnets",
                "Prompt Injection Defense: Untrusted context framing markers neutralize malicious instructions"
            ]
        },
        {
            "num": "11",
            "title": "Production Deployment Architecture",
            "subtitle": "Containerized & Cloud-Ready",
            "bullets": [
                "Multi-Container Docker Compose: Pre-configured services for Frontend, Backend, and PostgreSQL",
                "Static Frontend Hosting: Nginx Alpine container serving built React bundles with SPA fallback",
                "Stateless Backend + Persistent Volumes: Mounts for ChromaDB vectorstore and uploaded documents",
                "Automated Migrations: Alembic migrations execute automatically on container startup"
            ]
        },
        {
            "num": "12",
            "title": "Live Demonstration Walkthrough",
            "subtitle": "Validating the End-to-End Pipeline in Real-Time",
            "bullets": [
                "1. Submitting inquiry: 'Compare LangGraph, CrewAI, and PydanticAI for production AI'",
                "2. Live Workspace: Observing the Planner, Researcher, and Verifier timeline events via SSE",
                "3. Source Discovery: Examining live captured source cards and confidence ratings",
                "4. Reading Report: Reviewing comparison tables, recommendations, and verified citations",
                "5. Export: Downloading markdown dossier and validating reference integrity"
            ]
        },
        {
            "num": "13",
            "title": "Limitations & Lessons Learned",
            "subtitle": "Transparent System Boundaries",
            "bullets": [
                "Paywall & Captcha Boundaries: Authenticated or bot-blocked sites cannot be crawled autonomously",
                "Scanned Document Handling: Pure raster scans without OCR text layers are cleanly rejected",
                "Consensus Divergence: Highly disputed technical or political queries highlight divergence rather than picking winners",
                "Cost & Latency Bounds: Hard caps on iterations prevent runaway LLM API consumption"
            ]
        },
        {
            "num": "14",
            "title": "Future Roadmap",
            "subtitle": "Scaling Multi-Agent Research Capabilities",
            "bullets": [
                "Headless Browser Crawling: Integrate Playwright workers for client-side JavaScript applications",
                "Distributed Task Queues: Scale beyond in-memory asyncio to Redis + Celery/Temporal workers",
                "Hybrid Dense-Sparse RAG: Combine BM25 lexical search with dense vector embeddings",
                "Interactive Human-in-the-Loop: Allow users to review and steer the research plan before research starts"
            ]
        },
        {
            "num": "15",
            "title": "Conclusion & Showcase Summary",
            "subtitle": "A Commercial-Ready Blueprint for Verifiable AI Research",
            "bullets": [
                "Complete Full-Stack Application built from scratch with clean architectural separation",
                "100% Verified Multi-Agent Orchestration via LangGraph StateGraph",
                "Authentic Evidence Grounding with zero hallucinated citations",
                "Production-tested, containerized, documented, and ready for deployment"
            ]
        }
    ]

    blank_slide_layout = prs.slide_layouts[6]

    for slide_data in slides_content:
        slide = prs.slides.add_slide(blank_slide_layout)

        # Background card
        bg_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg_shape.fill.solid()
        bg_shape.fill.fore_color.rgb = COLOR_BG
        bg_shape.line.fill.background()

        # Inner Panel
        panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.8), Inches(11.733), Inches(5.9))
        panel.fill.solid()
        panel.fill.fore_color.rgb = COLOR_CARD
        panel.line.color.rgb = COLOR_BORDER
        panel.line.width = Pt(1.5)

        # Top Accent Header Bar
        top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.8), Inches(11.733), Inches(0.12))
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = COLOR_PRIMARY
        top_bar.line.fill.background()

        # Slide Number Badge
        badge = slide.shapes.add_textbox(Inches(11.2), Inches(1.1), Inches(1.0), Inches(0.5))
        tf_b = badge.text_frame
        p_b = tf_b.paragraphs[0]
        p_b.text = slide_data["num"]
        p_b.font.size = Pt(14)
        p_b.font.bold = True
        p_b.font.color.rgb = COLOR_ACCENT
        p_b.alignment = PP_ALIGN.RIGHT

        # Title & Subtitle Box
        title_box = slide.shapes.add_textbox(Inches(1.2), Inches(1.1), Inches(9.8), Inches(1.3))
        tf = title_box.text_frame
        tf.word_wrap = True

        p_title = tf.paragraphs[0]
        p_title.text = slide_data["title"]
        p_title.font.size = Pt(28)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_TEXT

        p_sub = tf.add_paragraph()
        p_sub.text = slide_data["subtitle"]
        p_sub.font.size = Pt(14)
        p_sub.font.color.rgb = COLOR_MUTED
        p_sub.space_before = Pt(4)

        # Content Box
        content_box = slide.shapes.add_textbox(Inches(1.2), Inches(2.6), Inches(10.8), Inches(3.6))
        tf_c = content_box.text_frame
        tf_c.word_wrap = True

        for i, bullet in enumerate(slide_data["bullets"]):
            p = tf_c.paragraphs[0] if i == 0 else tf_c.add_paragraph()
            p.text = f"•  {bullet}"
            p.font.size = Pt(16)
            p.font.color.rgb = COLOR_TEXT
            p.space_after = Pt(16)

    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    output_path = os.path.join(base_dir, "docs", "Deep_Research_AI_Showcase.pptx")
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    prs.save(output_path)
    print(f"Presentation saved successfully to {output_path}")

if __name__ == "__main__":
    create_deck()
